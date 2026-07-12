"""
resume_parser.py
Extracts raw text from uploaded resume files (PDF or DOCX).
"""
import os
from PyPDF2 import PdfReader
from docx import Document


class ResumeParseError(Exception):
    """Raised when a resume file cannot be parsed."""


def extract_text_from_pdf(filepath: str) -> str:
    try:
        reader = PdfReader(filepath)
        text_chunks = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text_chunks.append(page_text)
        text = "\n".join(text_chunks).strip()
        if not text:
            raise ResumeParseError(
                "No extractable text found in PDF. It may be a scanned image."
            )
        return text
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(f"Failed to read PDF file: {exc}") from exc


def extract_text_from_docx(filepath: str) -> str:
    try:
        doc = Document(filepath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also pull text from tables (common in resumes with skill tables)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        text = "\n".join(paragraphs).strip()
        if not text:
            raise ResumeParseError("No extractable text found in DOCX file.")
        return text
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(f"Failed to read DOCX file: {exc}") from exc


def extract_resume_text(filepath: str) -> str:
    """Dispatch to the correct extractor based on file extension."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(filepath)
    else:
        raise ResumeParseError(f"Unsupported file type: {ext}")


def allowed_file(filename: str, allowed_extensions=("pdf", "docx", "doc")) -> bool:
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in allowed_extensions
    )
